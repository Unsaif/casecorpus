"""Human verification of extracted records: review workbooks out, decisions and metrics in.

`export_review` flattens every record into facts (one row per extracted item: value, quote, anchor, verification,
grounding) and writes one workbook per reviewer batch:
  README        instructions, legend, verdict definitions, progress counters (formulas)
  Facts         one row per fact; reviewer fills Verdict / Correction / Comment (yellow cells)
  Record        one row per record: diagnosis, narrative completeness, overall quality
  Missed        free rows for facts the paper states but the record lacks (recall)
  Paper text    every anchored unit of each paper as the extractor saw it, with the records whose
                narrative includes the unit — so quotes can be checked without opening the PDF
`import_review` reads filled workbooks back into gold/decisions.jsonl and computes per-field precision, recall,
grounding accuracy, attribution errors and inter-reviewer agreement.
"""
from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from .anchors import index_markdown
from .config import Settings
from .db import Catalogue

VERDICTS = ["Correct", "Partly correct", "Wrong value", "Wrong grounding", "Wrong patient", "Not in paper", "Unsure"]
YESNO = ["Yes", "No", "Partly"]
QUALITY = ["1", "2", "3", "4", "5"]
GROUPS = ["subject", "family", "phenotype", "measurement", "enzyme_activity", "genetic_finding", "diagnosis", "treatment", "outcome"]
GROUNDED_GROUPS = {"phenotype", "measurement", "genetic_finding", "diagnosis"}

FONT = "Arial"


# ------------------------------------------------------------------------------------------ facts
def _ev(item: dict[str, Any] | None) -> tuple[str, str, str]:
    ev = (item or {}).get("evidence") or {}
    v = ev.get("verification") or {}
    if not ev:
        return "", "", ""
    if v:
        status = "in anchor" if v.get("in_anchor") else ("in document" if v.get("in_document") else "NOT FOUND")
    else:
        status = "unverified"
    return ev.get("quote") or "", ev.get("anchor") or "", status


def _age(a: dict[str, Any] | None) -> str:
    if not a:
        return ""
    return " / ".join(x for x in (a.get("iso8601"), a.get("text")) if x)


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (list, tuple)):
        return "; ".join(_fmt(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def facts_of(rec: dict[str, Any], record_id: str, pmid: str) -> list[dict[str, Any]]:
    """Flatten a record into reviewable facts with stable ids (record_id#group:index)."""
    facts: list[dict[str, Any]] = []
    local_id = rec.get("local_id") or ""

    def add(group: str, i: int, item: str, detail: str, grounded: str, ev_from: dict[str, Any] | None) -> None:
        q, a, s = _ev(ev_from)
        facts.append({"fact_id": f"{record_id}#{group}:{i}", "record_id": record_id, "pmid": pmid, "local_id": local_id,
                      "group": group, "item": item, "detail": detail, "grounded": grounded, "quote": q, "anchor": a, "verification": s})

    subj = rec.get("subject") or {}
    i = 0
    for key in ("sex", "vital_status", "ancestry_as_stated", "country_of_care", "gestational_age_weeks", "birth_weight_g"):
        if subj.get(key) not in (None, "", "UNKNOWN_SEX", "UNKNOWN"):
            add("subject", i, key, _fmt(subj[key]), "", None); i += 1
    for key in ("age_at_onset", "age_at_presentation", "age_at_last_followup", "age_at_death"):
        a = subj.get(key)
        if a and (a.get("iso8601") or a.get("text")):
            add("subject", i, key, _age(a), "", a); i += 1
    fam = rec.get("family") or {}
    i = 0
    if fam.get("consanguinity") is not None:
        add("family", i, "consanguinity", _fmt(fam["consanguinity"]), "", None); i += 1
    if fam.get("family_history_text"):
        add("family", i, "family history", fam["family_history_text"], "", None); i += 1
    for i, p in enumerate(rec.get("phenotypes") or []):
        item = ("ABSENT: " if p.get("excluded") else "") + (p.get("text") or "")
        detail = "; ".join(x for x in (f"onset {_age(p.get('onset'))}" if _age(p.get("onset")) else "", p.get("severity") or "", p.get("modifier") or "") if x)
        g = " ".join(x for x in (p.get("hpo_id") or "", p.get("hpo_label") or "") if x)
        add("phenotype", i, item, detail, g, p)
    for i, m in enumerate(rec.get("measurements") or []):
        val = m.get("value_text") or (_fmt(m.get("value")) + (" " + m.get("unit") if m.get("unit") else ""))
        if m.get("comparator") and m.get("value") is not None and not m.get("value_text"):
            val = f"{m['comparator']}{val}"
        detail = "; ".join(x for x in (val, f"ref {m['reference_range']}" if m.get("reference_range") else "",
                                        m.get("interpretation") if m.get("interpretation") not in (None, "UNKNOWN") else "",
                                        f"specimen {m['specimen']}" if m.get("specimen") not in (None, "other") else "",
                                        f"timing {m['timing']}" if m.get("timing") else "", f"age {_age(m.get('age'))}" if _age(m.get("age")) else "") if x)
        an = m.get("analyte") or {}
        g = " ".join(x for x in (an.get("chebi") or "", an.get("label") or "") if x)
        add("measurement", i, m.get("analyte_text") or "", detail, g, m)
    for i, e in enumerate(rec.get("enzyme_activities") or []):
        detail = "; ".join(x for x in (f"{_fmt(e.get('value'))} {e.get('unit') or ''}".strip() if e.get("value") is not None else "",
                                        f"{_fmt(e.get('percent_of_control'))}% of control" if e.get("percent_of_control") is not None else "",
                                        f"tissue {e['tissue']}" if e.get("tissue") else "", e.get("interpretation") or "", e.get("reference_text") or "") if x)
        add("enzyme_activity", i, e.get("enzyme_text") or "", detail, " ".join(x for x in (e.get("gene_symbol") or "", e.get("ec_number") or "") if x), e)
    for i, v in enumerate(rec.get("genetic_findings") or []):
        item = " ".join(x for x in (v.get("gene_symbol") or "", v.get("variant_hgvs_c") or "", v.get("variant_hgvs_p") or "") if x) or (v.get("variant_text") or "")
        detail = "; ".join(x for x in (v.get("zygosity") or "", v.get("inheritance_as_stated") or "", v.get("classification_as_reported") or "",
                                        v.get("method") or "", f"transcript {v['transcript']}" if v.get("transcript") else "", "novel" if v.get("novel") else "") if x)
        add("genetic_finding", i, item, detail, v.get("hgnc_id") or "", v)
    for i, d in enumerate(rec.get("diagnoses") or []):
        dis = d.get("disease") or {}
        detail = "; ".join(x for x in (d.get("status") or "", d.get("confidence_tier") or "", d.get("diagnostic_route") or "",
                                        f"age at dx {_age(d.get('age_at_diagnosis'))}" if _age(d.get("age_at_diagnosis")) else "",
                                        f"prior: {_fmt(d['prior_misdiagnoses'])}" if d.get("prior_misdiagnoses") else "") if x)
        g = " ".join(x for x in (dis.get("mondo") or "", dis.get("label") or "") if x)
        add("diagnosis", i, d.get("disease_text") or "", detail, g, d)
    for i, t in enumerate(rec.get("treatments") or []):
        detail = "; ".join(x for x in (t.get("type") or "", t.get("agent") or "", f"response: {t['response']}" if t.get("response") else "", t.get("response_text") or "") if x)
        add("treatment", i, t.get("text") or "", detail, "", t)
    if (rec.get("outcome") or {}).get("text"):
        add("outcome", 0, "outcome", rec["outcome"]["text"], "", rec["outcome"])
    return facts


# ------------------------------------------------------------------------------------------ export
def _units(md: str) -> list[dict[str, str]]:
    """Every anchored unit of a rendered document, in order."""
    idx = index_markdown(md)
    out = [{"anchor": "title", "section": "", "text": idx.title}] if idx.title else []
    if idx.abstract:
        out.append({"anchor": "abstract", "section": "Abstract", "text": idx.abstract})
    for n in sorted(idx.sections):
        sec = idx.sections[n]
        for m in sorted(sec["paragraphs"]):
            out.append({"anchor": f"sec:{n} ¶{m}", "section": sec.get("title", ""), "text": sec["paragraphs"][m]})
    for n in sorted(idx.tables):
        tab = idx.tables[n]
        out.append({"anchor": f"tab:{n}", "section": f"Table {tab.get('label', n)}", "text": tab.get("caption", "")})
        for i in sorted(tab["rows"]):
            out.append({"anchor": f"tab:{n} r{i}", "section": f"Table {tab.get('label', n)}", "text": " | ".join(tab["rows"][i])})
    for k, v in idx.figures.items():
        out.append({"anchor": f"fig:{k}", "section": "Figure", "text": v})
    return out


def assign(record_ids: list[str], reviewers: list[str], batch_size: int = 8, double_fraction: float = 0.25, seed: int = 1) -> list[dict[str, Any]]:
    """Deterministic assignment: records shuffled by hash, dealt round-robin to reviewers in batches; the first
    double_fraction of each batch is also given to the next reviewer (for agreement)."""
    ids = sorted(record_ids, key=lambda r: hashlib.sha1(f"{seed}:{r}".encode()).hexdigest())
    out = []
    n_rev = len(reviewers)
    batches: dict[str, list[list[str]]] = {r: [[]] for r in reviewers}
    for k, rid in enumerate(ids):
        r = reviewers[k % n_rev]
        if len(batches[r][-1]) >= batch_size:
            batches[r].append([])
        batches[r][-1].append(rid)
    for r in reviewers:
        for b, ids_b in enumerate(batches[r], 1):
            for j, rid in enumerate(ids_b):
                out.append({"record_id": rid, "reviewer": r, "batch": b, "role": "primary"})
                if n_rev > 1 and j < round(len(ids_b) * double_fraction):
                    other = reviewers[(reviewers.index(r) + 1) % n_rev]
                    out.append({"record_id": rid, "reviewer": other, "batch": f"{b}x", "role": "second"})
    return out


def export_review(settings: Settings, cat: Catalogue, out_dir: Path, record_ids: list[str] | None = None, only_set: str | None = None,
                  reviewers: list[str] | None = None, batch_size: int = 8, double_fraction: float = 0.25, seed: int = 1) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    out_dir.mkdir(parents=True, exist_ok=True)
    reviewers = reviewers or ["reviewer1"]
    where, params = "1=1", ()
    if record_ids:
        where = "record_id IN (" + ",".join("?" for _ in record_ids) + ")"; params = tuple(record_ids)
    elif only_set:
        where = "pmid IN (SELECT pmid FROM document_sets WHERE set_name=?)"; params = (only_set,)
    recs = {r["record_id"]: r for r in cat.iter_individuals(where, params)}
    if not recs:
        return {"workbooks": 0, "records": 0, "error": "no records matched"}
    plan = assign(sorted(recs), reviewers, batch_size, double_fraction, seed)
    with open(out_dir / "assignments.tsv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["record_id", "reviewer", "batch", "role"], delimiter="\t"); w.writeheader(); w.writerows(plan)

    hdr_fill = PatternFill("solid", fgColor="DDE4EE")
    edit_fill = PatternFill("solid", fgColor="FFF7C2")
    bold = Font(name=FONT, bold=True)
    plain = Font(name=FONT)
    note = Font(name=FONT, italic=True, size=9, color="555555")

    def header(ws, cols: list[str], widths: dict[str, float] | None = None) -> None:
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c); cell.font = bold; cell.fill = hdr_fill; cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        for k, v in (widths or {}).items():
            ws.column_dimensions[k].width = v

    groups: dict[tuple[str, Any], list[str]] = defaultdict(list)
    for a in plan:
        groups[(a["reviewer"], a["batch"])].append(a["record_id"])
    n_books = 0
    for (rev, batch), rids in sorted(groups.items(), key=lambda kv: (kv[0][0], str(kv[0][1]))):
        wb = Workbook()
        # ---------------- README
        ws = wb.active; ws.title = "README"
        lines = [
            ("casecorpus review workbook", ""), ("Reviewer", rev), ("Batch", str(batch)), ("Records in this batch", str(len(rids))),
            ("Your name (fill in)", ""), ("Date (fill in)", ""), ("", ""),
            ("How to review", "Work record by record. Filter the Facts sheet on record_id. For each row read the extracted item, then check it against the quote and its place in the paper (Paper text sheet, same anchors; or the PDF). Fill Verdict; add a Correction when the value is wrong; Comment is free text. Yellow cells are yours; do not edit the others."),
            ("Verdicts", ""),
            ("Correct", "The item is stated in the paper for THIS individual, value and qualifier as extracted."),
            ("Partly correct", "Right item, but a detail is off (wrong unit, wrong onset, incomplete qualifier). Put the right detail in Correction."),
            ("Wrong value", "Item exists but the value/qualifier is wrong (e.g. 320 vs 32.0 µmol/L; 'absent' vs present). Put the right value in Correction."),
            ("Wrong grounding", "Text is right but the ontology id/label (HPO, MONDO, ChEBI, HGNC) is wrong. Put the right id or label in Correction."),
            ("Wrong patient", "Stated in the paper, but about another individual (series/family papers)."),
            ("Not in paper", "Not stated anywhere in the paper (invented, or textbook knowledge imported)."),
            ("Unsure", "Cannot decide from the paper; explain in Comment."),
            ("", ""),
            ("Missed sheet", "Facts the paper states about the individual that are NOT in the record: one row each — record_id, group, what the paper says, where (anchor or page). This is how we measure recall, so be complete for the groups that matter (phenotype, measurement, genetic_finding, diagnosis)."),
            ("Record sheet", "One row per record: is the FINAL diagnosis right, is its confidence tier right, does the verbatim narrative cover the whole description of this individual, does it include text about other individuals, overall quality 1-5."),
            ("Paper text", "The document exactly as the extractor saw it, one row per anchored unit (paragraph / table row). The 'narrative of' column lists the records whose verbatim narrative includes that unit."),
            ("Example (Facts row)", "phenotype | Hyperammonemia | onset P3D | HP:0001987 Hyperammonemia | quote 'plasma ammonia was 480 µmol/L on day 3' | sec:2 ¶2 | in anchor  ->  Verdict: Correct"),
        ]
        for a, b in lines:
            ws.append([a, b])
        for r in range(1, len(lines) + 1):
            ws.cell(row=r, column=1).font = bold if ws.cell(row=r, column=1).value in ("casecorpus review workbook", "Verdicts", "How to review") or r <= 6 else plain
            ws.cell(row=r, column=2).font = plain
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        for r in (5, 6):
            ws.cell(row=r, column=2).fill = edit_fill
        ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 120
        # ---------------- Facts
        wf = wb.create_sheet("Facts")
        cols = ["fact_id", "record_id", "pmid", "local_id", "group", "item", "detail", "grounded", "quote", "anchor", "quote verification", "Verdict", "Correction", "Comment"]
        header(wf, cols, {"A": 30, "B": 26, "C": 10, "D": 12, "E": 14, "F": 34, "G": 40, "H": 30, "I": 60, "J": 12, "K": 14, "L": 16, "M": 30, "N": 30})
        n_facts = 0
        for rid in rids:
            r = recs[rid]
            for f in facts_of(r["record"], rid, r["pmid"]):
                wf.append([f["fact_id"], f["record_id"], f["pmid"], f["local_id"], f["group"], f["item"], f["detail"], f["grounded"], f["quote"], f["anchor"], f["verification"], None, None, None])
                n_facts += 1
        last = n_facts + 1
        for row in wf.iter_rows(min_row=2, max_row=last):
            for c in row:
                c.font = plain
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for c in row[11:14]:
                c.fill = edit_fill
        dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
        wf.add_data_validation(dv); dv.add(f"L2:L{max(last, 2)}")
        wf.auto_filter.ref = f"A1:N{max(last, 2)}"
        # ---------------- Record
        wr = wb.create_sheet("Record")
        rcols = ["record_id", "pmid", "local_id", "extracted FINAL diagnosis", "tier", "n facts", "Diagnosis correct?", "Tier correct?", "Narrative complete?", "Narrative has other individuals' text?", "Overall quality (1-5)", "Comment"]
        header(wr, rcols, {"A": 26, "B": 10, "C": 12, "D": 40, "E": 12, "F": 8, "G": 16, "H": 12, "I": 16, "J": 20, "K": 14, "L": 50})
        for k, rid in enumerate(rids, 2):
            r = recs[rid]; rec = r["record"]
            fin = [d for d in rec.get("diagnoses") or [] if d.get("status") == "FINAL"] or (rec.get("diagnoses") or [])
            dx = "; ".join(f"{d.get('disease_text')} [{(d.get('disease') or {}).get('mondo') or '?'}]" for d in fin[:2])
            tier = "; ".join(d.get("confidence_tier") or "" for d in fin[:2])
            wr.append([rid, r["pmid"], rec.get("local_id"), dx, tier, f'=COUNTIF(Facts!B2:B{max(last, 2)},A{k})', None, None, None, None, None, None])
            for c in wr[k]:
                c.font = plain; c.alignment = Alignment(wrap_text=True, vertical="top")
            for c in wr[k][6:12]:
                c.fill = edit_fill
        rl = len(rids) + 1
        dv1 = DataValidation(type="list", formula1='"Yes,No,Partly"', allow_blank=True); wr.add_data_validation(dv1); dv1.add(f"G2:J{max(rl, 2)}")
        dv2 = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True); wr.add_data_validation(dv2); dv2.add(f"K2:K{max(rl, 2)}")
        # ---------------- Missed
        wm = wb.create_sheet("Missed")
        header(wm, ["record_id", "group", "What the paper states (the missed fact)", "Where (anchor or page)", "Comment"], {"A": 26, "B": 16, "C": 60, "D": 18, "E": 40})
        for i in range(2, 402):
            for c in range(1, 6):
                cell = wm.cell(row=i, column=c); cell.font = plain; cell.fill = edit_fill; cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv3 = DataValidation(type="list", formula1=f"=Record!$A$2:$A${max(rl, 2)}", allow_blank=True); wm.add_data_validation(dv3); dv3.add("A2:A401")
        dv4 = DataValidation(type="list", formula1='"' + ",".join(GROUPS) + '"', allow_blank=True); wm.add_data_validation(dv4); dv4.add("B2:B401")
        # ---------------- Paper text
        wp = wb.create_sheet("Paper text")
        header(wp, ["pmid", "anchor", "section", "text", "narrative of"], {"A": 10, "B": 12, "C": 22, "D": 120, "E": 26})
        seen: set[str] = set()
        for rid in rids:
            pmid = recs[rid]["pmid"]
            if pmid in seen:
                continue
            seen.add(pmid)
            md_path = settings.work_dir / pmid / "input.md"
            md = md_path.read_text() if md_path.exists() else ""
            narrative_of: dict[str, list[str]] = defaultdict(list)
            for r2 in recs.values():
                if r2["pmid"] == pmid:
                    for u in r2["record"].get("case_narrative") or []:
                        narrative_of[u.get("anchor")].append(r2["record"].get("local_id") or r2["record_id"])
            for u in _units(md):
                wp.append([pmid, u["anchor"], u["section"], u["text"], ", ".join(narrative_of.get(u["anchor"], []))])
        for row in wp.iter_rows(min_row=2):
            for c in row:
                c.font = plain; c.alignment = Alignment(wrap_text=True, vertical="top")
        wp.auto_filter.ref = f"A1:E{max(wp.max_row, 2)}"
        # ---------------- progress counters on README (formulas)
        base = len(lines) + 2
        ws.cell(row=base, column=1, value="Progress").font = bold
        ws.cell(row=base + 1, column=1, value="Facts in batch").font = plain
        ws.cell(row=base + 1, column=2, value=f"=COUNTA(Facts!A2:A{max(last, 2)})").font = plain
        ws.cell(row=base + 2, column=1, value="Facts with a verdict").font = plain
        ws.cell(row=base + 2, column=2, value=f"=COUNTA(Facts!L2:L{max(last, 2)})").font = plain
        ws.cell(row=base + 3, column=1, value="Records with a diagnosis verdict").font = plain
        ws.cell(row=base + 3, column=2, value=f"=COUNTA(Record!G2:G{max(rl, 2)})").font = plain
        ws.cell(row=base + 4, column=1, value="Missed facts entered").font = plain
        ws.cell(row=base + 4, column=2, value="=COUNTA(Missed!C2:C401)").font = plain
        ws.cell(row=base + 5, column=1, value="Progress counters are formulas over the other sheets; everything else on this sheet is text.").font = note
        wb.save(out_dir / f"review_{rev}_{batch}.xlsx")
        n_books += 1
    return {"workbooks": n_books, "records": len(recs), "assignments": len(plan), "reviewers": reviewers, "dir": str(out_dir)}


# ------------------------------------------------------------------------------------------ import
def _rows(ws) -> list[dict[str, Any]]:
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(it, [])]
    out = []
    for row in it:
        if row is None or all(v in (None, "") for v in row):
            continue
        out.append({hdr[i]: row[i] for i in range(min(len(hdr), len(row)))})
    return out


def import_review(files: Iterable[Path], gold_dir: Path) -> dict[str, Any]:
    """Read filled review workbooks; write gold/decisions.jsonl, gold/missed.jsonl, gold/records.jsonl and metrics."""
    from openpyxl import load_workbook
    gold_dir.mkdir(parents=True, exist_ok=True)
    decisions: list[dict[str, Any]] = []
    missed: list[dict[str, Any]] = []
    record_rows: list[dict[str, Any]] = []
    for f in files:
        wb = load_workbook(f, read_only=True, data_only=True)
        readme = {str(r[0]): r[1] for r in wb["README"].iter_rows(values_only=True) if r and r[0]}
        reviewer = readme.get("Your name (fill in)") or readme.get("Reviewer") or f.stem
        for r in _rows(wb["Facts"]):
            if r.get("Verdict"):
                decisions.append({"reviewer": str(reviewer), "file": f.name, **{k: r.get(k) for k in ("fact_id", "record_id", "pmid", "local_id", "group", "item", "detail", "grounded", "anchor")},
                                  "verdict": r["Verdict"], "correction": r.get("Correction"), "comment": r.get("Comment")})
        for r in _rows(wb["Missed"]):
            if r.get("record_id") and r.get("What the paper states (the missed fact)"):
                missed.append({"reviewer": str(reviewer), "file": f.name, "record_id": r["record_id"], "group": r.get("group") or "unknown",
                               "text": r["What the paper states (the missed fact)"], "where": r.get("Where (anchor or page)"), "comment": r.get("Comment")})
        for r in _rows(wb["Record"]):
            if any(r.get(k) for k in ("Diagnosis correct?", "Tier correct?", "Narrative complete?", "Overall quality (1-5)")):
                record_rows.append({"reviewer": str(reviewer), "file": f.name, "record_id": r.get("record_id"), "pmid": r.get("pmid"),
                                    "diagnosis_correct": r.get("Diagnosis correct?"), "tier_correct": r.get("Tier correct?"),
                                    "narrative_complete": r.get("Narrative complete?"), "narrative_has_others": r.get("Narrative has other individuals' text?"),
                                    "quality": r.get("Overall quality (1-5)"), "comment": r.get("Comment")})
    for name, rows in (("decisions.jsonl", decisions), ("missed.jsonl", missed), ("records.jsonl", record_rows)):
        with open(gold_dir / name, "w") as fh:
            for r in rows:
                fh.write(json.dumps(r, ensure_ascii=False, default=str) + "\n")
    metrics = compute_metrics(decisions, missed, record_rows)
    (gold_dir / "metrics.json").write_text(json.dumps(metrics, indent=1))
    (gold_dir / "metrics.md").write_text(metrics_markdown(metrics))
    return metrics


def compute_metrics(decisions: list[dict[str, Any]], missed: list[dict[str, Any]], record_rows: list[dict[str, Any]]) -> dict[str, Any]:
    # one verdict per (fact, reviewer); for precision use the primary = first reviewer seen per fact, agreement uses pairs
    by_fact: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for d in decisions:
        by_fact[d["fact_id"]].append(d)
    reviewed_records = {d["record_id"] for d in decisions}
    per_group: dict[str, Counter] = defaultdict(Counter)
    for fid, ds in by_fact.items():
        d = ds[0]
        per_group[d["group"]][d["verdict"]] += 1
        per_group["all"][d["verdict"]] += 1
    missed_by_group: Counter = Counter()
    for m in missed:
        if m["record_id"] in reviewed_records or not reviewed_records:
            missed_by_group[m["group"]] += 1
            missed_by_group["all"] += 1

    def rates(c: Counter, group: str) -> dict[str, Any]:
        n = sum(v for k, v in c.items() if k != "Unsure")
        correct = c["Correct"]; partly = c["Partly correct"]
        miss = missed_by_group.get(group, 0)
        return {"reviewed": n, "unsure": c["Unsure"], "correct": correct, "partly_correct": partly, "wrong_value": c["Wrong value"],
                "wrong_grounding": c["Wrong grounding"], "wrong_patient": c["Wrong patient"], "not_in_paper": c["Not in paper"], "missed": miss,
                "precision_strict": round(correct / n, 3) if n else None, "precision_lenient": round((correct + partly) / n, 3) if n else None,
                "recall_strict": round(correct / (correct + partly + miss), 3) if (correct + partly + miss) else None,
                "recall_lenient": round((correct + partly) / (correct + partly + miss), 3) if (correct + partly + miss) else None,
                "grounding_error_rate": round(c["Wrong grounding"] / n, 3) if (n and group in GROUNDED_GROUPS | {"all"}) else None,
                "attribution_error_rate": round(c["Wrong patient"] / n, 3) if n else None}

    out: dict[str, Any] = {"n_decisions": len(decisions), "n_facts_reviewed": len(by_fact), "n_records_reviewed": len(reviewed_records),
                           "n_missed": len(missed), "by_group": {g: rates(per_group[g], g) for g in ["all"] + [g for g in GROUPS if g in per_group]}}
    # agreement on double-reviewed facts (Correct vs not-Correct, and exact verdict)
    pairs = [(ds[0]["verdict"], ds[1]["verdict"]) for ds in by_fact.values() if len(ds) >= 2 and ds[0]["reviewer"] != ds[1]["reviewer"]]
    if pairs:
        a = sum(1 for x, y in pairs if x == y) / len(pairs)
        bx = [(x == "Correct", y == "Correct") for x, y in pairs]
        po = sum(1 for x, y in bx if x == y) / len(bx)
        p1 = sum(1 for x, _ in bx if x) / len(bx); p2 = sum(1 for _, y in bx if y) / len(bx)
        pe = p1 * p2 + (1 - p1) * (1 - p2)
        kappa = (po - pe) / (1 - pe) if pe < 1 else None
        out["agreement"] = {"double_reviewed_facts": len(pairs), "exact_verdict_agreement": round(a, 3), "correct_vs_not_agreement": round(po, 3),
                            "cohen_kappa_correct_vs_not": round(kappa, 3) if kappa is not None else None}
    else:
        out["agreement"] = {"double_reviewed_facts": 0}
    if record_rows:
        rc = Counter(r["diagnosis_correct"] for r in record_rows if r.get("diagnosis_correct"))
        tc = Counter(r["tier_correct"] for r in record_rows if r.get("tier_correct"))
        nc = Counter(r["narrative_complete"] for r in record_rows if r.get("narrative_complete"))
        no = Counter(r["narrative_has_others"] for r in record_rows if r.get("narrative_has_others"))
        q = [float(r["quality"]) for r in record_rows if r.get("quality") not in (None, "")]
        out["records"] = {"n": len(record_rows), "diagnosis_correct": dict(rc), "tier_correct": dict(tc), "narrative_complete": dict(nc),
                          "narrative_has_other_individuals": dict(no), "mean_quality": round(sum(q) / len(q), 2) if q else None}
    return out


def metrics_markdown(m: dict[str, Any]) -> str:
    lines = ["# Review metrics", "", f"Decisions {m['n_decisions']} on {m['n_facts_reviewed']} facts in {m['n_records_reviewed']} records; {m['n_missed']} missed facts entered.", "",
             "| group | reviewed | correct | partly | wrong value | wrong grounding | wrong patient | not in paper | missed | precision (strict / lenient) | recall (strict / lenient) |",
             "|---|---|---|---|---|---|---|---|---|---|---|"]
    for g, r in m["by_group"].items():
        lines.append(f"| {g} | {r['reviewed']} | {r['correct']} | {r['partly_correct']} | {r['wrong_value']} | {r['wrong_grounding']} | {r['wrong_patient']} | {r['not_in_paper']} | {r['missed']} | {r['precision_strict']} / {r['precision_lenient']} | {r['recall_strict']} / {r['recall_lenient']} |")
    a = m.get("agreement", {})
    lines += ["", f"Agreement: {a.get('double_reviewed_facts', 0)} double-reviewed facts; exact-verdict agreement {a.get('exact_verdict_agreement')}, correct-vs-not agreement {a.get('correct_vs_not_agreement')}, Cohen's kappa {a.get('cohen_kappa_correct_vs_not')}."]
    if m.get("records"):
        r = m["records"]
        lines += ["", f"Record level (n={r['n']}): diagnosis correct {r['diagnosis_correct']}; tier correct {r['tier_correct']}; narrative complete {r['narrative_complete']}; narrative has other individuals' text {r['narrative_has_other_individuals']}; mean quality {r['mean_quality']}."]
    return "\n".join(lines) + "\n"
