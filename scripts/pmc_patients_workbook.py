"""Build the PMC-Patients overview workbook from the overview directory produced by
`casecorpus pmc-patients-overview`. Aggregates that can be expressed as COUNTIFS/PERCENTILE over the
Patients sheet are formulas; distinct-article counts and the disease frequency tables are values computed
in Python and labelled as such.

Usage: python pmc_patients_workbook.py <overview_dir> <out.xlsx>
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="DDE4EE")
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="555555")


def style_header(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{row + 1}"  # string form: ws.cell() would create an empty cell and shift the next append down a row


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def write_table(ws, df: pd.DataFrame, start_row: int = 1, note: str | None = None) -> int:
    """Write a DataFrame with a styled header; returns the next free row."""
    for j, col in enumerate(df.columns, 1):
        ws.cell(row=start_row, column=j, value=str(col))
    style_header(ws, len(df.columns), start_row)
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if pd.isna(v) if not isinstance(v, str) else False:
                v = None
            ws.cell(row=i, column=j, value=v).font = Font(name=FONT)
    nxt = start_row + len(df) + 1
    if note:
        ws.cell(row=nxt, column=1, value=note).font = NOTE_FONT
        nxt += 1
    return nxt + 1


def main(overview_dir: Path, out: Path) -> None:
    df = pd.read_parquet(overview_dir / "patients.parquet")
    stats = json.load(open(overview_dir / "overview.json"))
    df["year"] = df["approx_year"].round().astype("Int64")
    cols = ["patient_uid", "pmid", "year", "licence_group", "gender", "age_years", "age_band", "n_words", "patients_in_article",
            "iem", "iem_title", "iem_genes", "title_mondo_labels", "title_mondo", "title"]
    pat = df[cols].copy()
    pat["iem"] = pat["iem"].map({True: "yes", False: "no"})
    pat["iem_title"] = pat["iem_title"].map({True: "yes", False: "no"})
    n = len(pat)
    last = n + 1  # last data row on the Patients sheet

    wb = Workbook()
    # ---------------------------------------------------------------- Overview
    ov = wb.active
    ov.title = "Overview"
    P = "Patients"
    rows = [
        ("PMC-Patients V2 — overview", None, None),
        ("Source", "PMC-Patients V2 (Zhao et al., Sci Data 2023; V2 from the 2024 PMC baseline), zhengyun21/PMC-Patients on Hugging Face, CC BY-NC-SA 4.0", None),
        ("Built", "casecorpus pmc-patients-overview, 11 Sept 2026", None),
        (None, None, None),
        ("Measure", "Value", "How"),
        ("Patient summaries", f"=COUNTA({P}!A2:A{last})", "COUNTA over the Patients sheet"),
        ("Source articles (distinct PMIDs)", stats["articles"], "computed in Python: distinct PMIDs (Excel has no simple distinct count)"),
        ("Female", f"=COUNTIF({P}!E2:E{last},\"F\")", "COUNTIF"),
        ("Male", f"=COUNTIF({P}!E2:E{last},\"M\")", "COUNTIF"),
        ("Median age (years)", f"=MEDIAN({P}!F2:F{last})", "MEDIAN"),
        ("Median summary length (words)", f"=MEDIAN({P}!H2:H{last})", "MEDIAN"),
        ("Commercial-use licence group (oa_comm)", f"=COUNTIF({P}!D2:D{last},\"comm\")", "COUNTIF on file path prefix"),
        ("Non-commercial licence group (oa_noncomm)", f"=COUNTIF({P}!D2:D{last},\"noncomm\")", "COUNTIF on file path prefix"),
        ("Titles naming at least one disease (MONDO match)", f"=COUNTIF({P}!M2:M{last},\"?*\")", "COUNTIF non-empty"),
        ("IEM-related (disease in MONDO IEM subtree named in title, or IEM gene symbol in text)", f"=COUNTIF({P}!J2:J{last},\"yes\")", "COUNTIF"),
        ("   of which IEM disease named in the title", f"=COUNTIF({P}!K2:K{last},\"yes\")", "COUNTIF"),
        ("Approximate year range", f"{int(df['year'].min())}–{int(df['year'].max())}", "from PMID (see Notes)"),
    ]
    for r, (a, b, c) in enumerate(rows, 1):
        for j, v in enumerate((a, b, c), 1):
            if v is not None:
                cell = ov.cell(row=r, column=j, value=v)
                cell.font = Font(name=FONT, bold=(r == 1 or r == 5), size=14 if r == 1 else 11)
                if j == 3:
                    cell.font = NOTE_FONT
    style_header(ov, 3, 5)
    ov.freeze_panes = None
    set_widths(ov, {"A": 62, "B": 40, "C": 60})
    for r in range(6, 6 + 12):
        ov.cell(row=r, column=2).number_format = "#,##0.0" if r in (10, 11) else "#,##0"

    # ---------------------------------------------------------------- Patients (full table)
    ws = wb.create_sheet("Patients")
    ws.append(list(pat.columns))
    style_header(ws, len(pat.columns))
    for row in pat.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) or v is pd.NA else v for v in row])
    set_widths(ws, {"A": 13, "B": 10, "C": 7, "D": 10, "E": 7, "F": 9, "G": 20, "H": 8, "I": 9, "J": 6, "K": 8, "L": 16, "M": 44, "N": 26, "O": 90})
    ws.auto_filter.ref = f"A1:{get_column_letter(len(pat.columns))}{last}"

    # ---------------------------------------------------------------- By year
    ws = wb.create_sheet("By year")
    years = list(range(int(df["year"].min()), int(df["year"].max()) + 1))
    art = df.groupby("year")["pmid"].nunique()
    ws.append(["approx_year", "patients", "articles (distinct PMIDs)"]); style_header(ws, 3)
    for i, y in enumerate(years, 2):
        ws.cell(row=i, column=1, value=y).font = Font(name=FONT)
        ws.cell(row=i, column=2, value=f"=COUNTIF({P}!C2:C{last},A{i})").font = Font(name=FONT)
        ws.cell(row=i, column=3, value=int(art.get(y, 0))).font = Font(name=FONT)
    r = len(years) + 2
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B2:B{r-1})").font = Font(name=FONT, bold=True)
    ws.cell(row=r + 1, column=1, value="Year is approximate: interpolated from the PMID (PubMed assigns PMIDs roughly in order of receipt), ±1 year. Article counts are distinct PMIDs computed in Python.").font = NOTE_FONT
    set_widths(ws, {"A": 12, "B": 12, "C": 24})

    # ---------------------------------------------------------------- Demographics
    ws = wb.create_sheet("Demographics")
    bands = ["neonate (<28 d)", "infant (<1 y)", "child (1-<12 y)", "adolescent (12-<18 y)", "adult (18-<65 y)", "older adult (>=65 y)", "unknown"]
    ws.append(["age_band", "female", "male", "all", "median_age_years"]); style_header(ws, 5)
    for i, b in enumerate(bands, 2):
        ws.cell(row=i, column=1, value=b)
        ws.cell(row=i, column=2, value=f"=COUNTIFS({P}!G2:G{last},A{i},{P}!E2:E{last},\"F\")")
        ws.cell(row=i, column=3, value=f"=COUNTIFS({P}!G2:G{last},A{i},{P}!E2:E{last},\"M\")")
        ws.cell(row=i, column=4, value=f"=COUNTIF({P}!G2:G{last},A{i})")
        med = df.loc[df["age_band"] == b, "age_years"].median()
        ws.cell(row=i, column=5, value=None if pd.isna(med) else round(float(med), 3))
        for c in range(1, 6):
            ws.cell(row=i, column=c).font = Font(name=FONT)
    r = len(bands) + 2
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT, bold=True)
    for c, col in ((2, "B"), (3, "C"), (4, "D")):
        ws.cell(row=r, column=c, value=f"=SUM({col}2:{col}{r-1})").font = Font(name=FONT, bold=True)
    ws.cell(row=r + 1, column=1, value="Age bands from the dataset's parsed age (value, unit) list collapsed to years. Median per band computed in Python.").font = NOTE_FONT
    r += 3
    ws.cell(row=r, column=1, value="licence_group").font = Font(name=FONT, bold=True); ws.cell(row=r, column=2, value="patients").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=3, value="articles (distinct PMIDs)").font = Font(name=FONT, bold=True)
    lic = df.groupby("licence_group")["pmid"].nunique()
    for k, g in enumerate(["comm", "noncomm", "other", "unknown"], r + 1):
        ws.cell(row=k, column=1, value=g).font = Font(name=FONT)
        ws.cell(row=k, column=2, value=f"=COUNTIF({P}!D2:D{last},A{k})").font = Font(name=FONT)
        ws.cell(row=k, column=3, value=int(lic.get(g, 0))).font = Font(name=FONT)
    ws.cell(row=r + 6, column=1, value="Licence group = PMC open-access package the source article sits in (oa_comm: CC BY-type, commercial reuse allowed; oa_noncomm: CC BY-NC-type). The dataset itself is CC BY-NC-SA 4.0.").font = NOTE_FONT
    set_widths(ws, {"A": 26, "B": 12, "C": 24, "D": 10, "E": 18})

    # ---------------------------------------------------------------- Articles & length
    ws = wb.create_sheet("Articles & length")
    ws.append(["patients_in_article", "articles", "patients"]); style_header(ws, 3)
    ks = sorted(df["patients_in_article"].unique())
    for i, k in enumerate(ks, 2):
        ws.cell(row=i, column=1, value=int(k)).font = Font(name=FONT)
        ws.cell(row=i, column=3, value=f"=COUNTIF({P}!I2:I{last},A{i})").font = Font(name=FONT)
        ws.cell(row=i, column=2, value=f"=C{i}/A{i}").font = Font(name=FONT)
    r = len(ks) + 2
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B2:B{r-1})").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=3, value=f"=SUM(C2:C{r-1})").font = Font(name=FONT, bold=True)
    r += 2
    ws.cell(row=r, column=1, value="Summary length (words)").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=2, value="value").font = Font(name=FONT, bold=True)
    for i, (lab, f) in enumerate([("minimum", f"=MIN({P}!H2:H{last})"), ("5th percentile", f"=PERCENTILE({P}!H2:H{last},0.05)"), ("25th percentile", f"=PERCENTILE({P}!H2:H{last},0.25)"),
                                  ("median", f"=MEDIAN({P}!H2:H{last})"), ("75th percentile", f"=PERCENTILE({P}!H2:H{last},0.75)"), ("95th percentile", f"=PERCENTILE({P}!H2:H{last},0.95)"),
                                  ("maximum", f"=MAX({P}!H2:H{last})"), ("mean", f"=AVERAGE({P}!H2:H{last})")], r + 1):
        ws.cell(row=i, column=1, value=lab).font = Font(name=FONT)
        ws.cell(row=i, column=2, value=f).font = Font(name=FONT)
        ws.cell(row=i, column=2).number_format = "#,##0.0"
    set_widths(ws, {"A": 26, "B": 14, "C": 12})

    # ---------------------------------------------------------------- Diseases in titles
    ws = wb.create_sheet("Diseases in titles")
    top = pd.read_csv(overview_dir / "top_title_diseases.csv")
    nxt = write_table(ws, top.rename(columns={"disease_in_title": "disease named in title (MONDO label)", "patients": "patients"}),
                      note="Computed in Python: exact match of MONDO disease labels/synonyms (normalised) against article titles; generic class names and qualifiers excluded. A title can name several diseases; counts are per patient summary. This is what the title says, not a verified diagnosis.")
    set_widths(ws, {"A": 48, "B": 12})

    # ---------------------------------------------------------------- IEM
    ws = wb.create_sheet("IEM")
    iem = pd.read_csv(overview_dir / "iem.csv")
    nxt = write_table(ws, iem, note="IEM = disease in the MONDO 'inborn errors of metabolism' subtree (casecorpus scope, 2,046 diseases) named in the title, or one of the scope's 1,225 gene symbols (>= 4 characters) in title or summary. The MONDO subtree is broader than the clinical IEM notion (it includes e.g. amyloidoses, porphyrias, periodic paralyses, hemochromatosis).")
    iem_top = pd.read_csv(overview_dir / "top_iem_title_diseases.csv")
    write_table(ws, iem_top.rename(columns={"disease_in_title": "IEM-subtree disease named in title", "patients": "patients"}), start_row=nxt)
    set_widths(ws, {"A": 48, "B": 12})

    # ---------------------------------------------------------------- Notes
    ws = wb.create_sheet("Notes")
    notes = [
        "What this is: one row per patient summary in PMC-Patients V2 (250,294 summaries cut from 210,069 open-access PMC case reports by the dataset authors' heuristics), with derived columns.",
        "Columns on the Patients sheet: patient_uid (PMID-index as in the dataset); pmid; year (approximate, from PMID); licence_group (oa_comm / oa_noncomm from the source file path); gender (as given); age_years (dataset age list collapsed to years); age_band; n_words (summary length); patients_in_article; iem / iem_title / iem_genes (see IEM sheet); title_mondo_labels and title_mondo (diseases named in the article title, MONDO); title.",
        "Not included here: the summary text itself (available in patients.parquet next to this workbook) and the dataset's relevant_articles / similar_patients annotations.",
        "Year caveat: PMC-Patients carries no publication date; year is interpolated from the PMID with a table of first-PMID-of-year anchors (±1 year).",
        "Disease caveat: title tagging is lexical (exact match of MONDO names/synonyms after normalisation). It finds what the title names; it does not know whether the named disease is the patient's diagnosis, a differential, or the topic of a series. Use the casecorpus extraction for diagnoses.",
        "Licence: PMC-Patients is CC BY-NC-SA 4.0 (non-commercial). Source articles in oa_noncomm are also non-commercial.",
        "Reproduce: casecorpus pmc-patients-overview PMC-Patients-V2.json.gz; python scripts/pmc_patients_workbook.py <overview_dir> <out.xlsx>",
    ]
    for i, t in enumerate(notes, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(ws, {"A": 140})

    wb.save(out)
    print("saved", out, "rows", n)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
